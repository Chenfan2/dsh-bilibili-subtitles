"""Word (.docx) renderer using python-docx."""

from __future__ import annotations

from pathlib import Path

from ..errors import DependencyMissingError, RenderError
from ..utils import ensure_dir, format_timestamp, human_duration, language_label
from .base import RenderContext, source_label

try:  # optional dependency, checked at call time
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except Exception:  # pragma: no cover - environment dependent
    Document = None  # type: ignore


EA_FONT = "微软雅黑"
LATIN_FONT = "Microsoft YaHei"
if Document is not None:
    GRAY = RGBColor(0x77, 0x77, 0x77)
    LIGHT_GRAY = RGBColor(0x99, 0x99, 0x99)
    ACCENT = RGBColor(0x1F, 0x4E, 0x79)
else:  # dependency is reported when rendering, not at import time
    GRAY = LIGHT_GRAY = ACCENT = None


def _set_run_font(run, *, size: float | None = None, bold: bool | None = None, color=None, italic=None):
    run.font.name = LATIN_FONT
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def _set_style_font(style, *, size: float, bold: bool = False, color=None):
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), EA_FONT)


def _add_page_number(paragraph) -> None:
    # A simple PAGE field; the cached result "1" is replaced by Word/LO.
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), " PAGE ")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Arial")
    rfonts.set(qn("w:hAnsi"), "Arial")
    rfonts.set(qn("w:eastAsia"), EA_FONT)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "16")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "999999")
    rpr.append(rfonts)
    rpr.append(size)
    rpr.append(color)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(rpr)
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def _add_footer(document, context: RenderContext) -> None:
    section = document.sections[0]
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("dsh-bilibili-subtitles · ")
    _set_run_font(run, size=8, color=LIGHT_GRAY)
    run = paragraph.add_run(context.video.title[:60])
    _set_run_font(run, size=8, color=LIGHT_GRAY)
    run = paragraph.add_run(" · ")
    _set_run_font(run, size=8, color=LIGHT_GRAY)
    _add_page_number(paragraph)


def _add_metadata_table(document, context: RenderContext, result) -> None:
    from .base import metadata_rows

    rows = metadata_rows(context, result)
    table = document.add_table(rows=0, cols=2)
    table.autofit = True
    try:
        table.style = "Light Grid Accent 1"
    except Exception:
        try:
            table.style = "Table Grid"
        except Exception:
            pass
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].width = Cm(3.1)
        cells[1].width = Cm(13.0)
        p0 = cells[0].paragraphs[0]
        run = p0.add_run(str(label))
        _set_run_font(run, size=9.5, bold=True)
        p1 = cells[1].paragraphs[0]
        run = p1.add_run(str(value))
        _set_run_font(run, size=9.5)
    document.add_paragraph()


def _add_quality(document, context: RenderContext) -> None:
    if not context.include_quality:
        return
    document.add_page_break()
    heading = document.add_heading("处理信息与质量报告", level=1)
    for run in heading.runs:
        _set_run_font(run, size=16, bold=True, color=ACCENT)
    quality = context.quality
    status = {"ok": "通过", "warn": "有警告", "fail": "失败"}.get(
        quality.get("status", ""), str(quality.get("status", "未知"))
    )
    lines = [
        f"状态：{status}    评分：{quality.get('score', '—')} / 100",
        f"时间覆盖率：{float(quality.get('coverage_ratio') or 0):.1%}    "
        f"片段：{quality.get('segments', 0)}   段落：{quality.get('paragraphs', 0)}",
        f"字符数：{quality.get('chars', 0)}    低置信度片段：{quality.get('low_confidence_segments', 0)}",
    ]
    for line in lines:
        p = document.add_paragraph()
        run = p.add_run(line)
        _set_run_font(run, size=9.5, color=GRAY)
    warnings = quality.get("warnings") or []
    if warnings:
        p = document.add_paragraph()
        run = p.add_run("警告")
        _set_run_font(run, size=10.5, bold=True)
        for warning in warnings:
            p = document.add_paragraph(style="List Bullet")
            run = p.add_run(warning)
            _set_run_font(run, size=9.5, color=GRAY)
    p = document.add_paragraph()
    run = p.add_run(
        "字幕由平台、UP 主或自动识别提供，可能存在错字、断句或时间轴误差；"
        "重要内容请对照原视频。请遵守版权与平台使用条款，仅处理你有权使用的视频。"
    )
    _set_run_font(run, size=8.5, color=LIGHT_GRAY, italic=True)


def render_docx(context: RenderContext, out_path: str | Path) -> Path:
    if Document is None:
        raise DependencyMissingError(
            "未安装 python-docx，无法生成 Word 文档。",
            hint="使用 DSH 自带 Python，或运行: python -m pip install python-docx",
        )
    out = Path(out_path)
    ensure_dir(out.parent)
    try:
        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.0)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.footer.is_linked_to_previous = False

        styles = document.styles
        normal = styles["Normal"]
        _set_style_font(normal, size=10.5)
        normal.paragraph_format.line_spacing = 1.45
        normal.paragraph_format.space_after = Pt(4)
        for name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
            try:
                _set_style_font(styles[name], size=size, bold=True, color=ACCENT)
            except Exception:
                pass

        title = document.add_heading(context.video.title, level=0)
        for run in title.runs:
            _set_run_font(run, size=22, bold=True, color=ACCENT)
        subtitle_parts = [context.video.owner or "", context.video.pubdate or "", human_duration(context.video.duration)]
        subtitle = "  ·  ".join(part for part in subtitle_parts if part)
        if subtitle:
            p = document.add_paragraph()
            run = p.add_run(subtitle)
            _set_run_font(run, size=10, color=GRAY)
        if context.video.url:
            p = document.add_paragraph()
            run = p.add_run(context.video.url)
            _set_run_font(run, size=9, color=GRAY)

        if context.parts:
            _add_metadata_table(document, context, context.parts[0])

        for result in context.parts:
            heading_text = (
                f"P{result.part.index} · {result.part.part_title or result.part.title}"
                if context.multi_part
                else "正文"
            )
            heading = document.add_heading(heading_text, level=1)
            for run in heading.runs:
                _set_run_font(run, size=16, bold=True, color=ACCENT)

            info = (
                f"字幕来源：{source_label(result)}    "
                f"语言：{language_label(result.track.language)}    "
                f"片段：{len(result.track.segments)}"
            )
            p = document.add_paragraph()
            run = p.add_run(info)
            _set_run_font(run, size=9, color=GRAY)

            if not result.paragraphs:
                p = document.add_paragraph()
                run = p.add_run("（没有可用字幕）")
                _set_run_font(run, size=10.5, italic=True, color=GRAY)
                continue

            for paragraph in result.paragraphs:
                p = document.add_paragraph()
                p.paragraph_format.space_after = Pt(6)
                if context.include_timestamp:
                    run = p.add_run(f"[{format_timestamp(paragraph.start)}] ")
                    _set_run_font(run, size=9, color=LIGHT_GRAY)
                if paragraph.speaker:
                    run = p.add_run(f"{paragraph.speaker}：")
                    _set_run_font(run, size=10.5, bold=True)
                run = p.add_run(paragraph.text)
                _set_run_font(run, size=10.5)
                if context.bilingual and paragraph.translation:
                    tp = document.add_paragraph()
                    tp.paragraph_format.left_indent = Cm(0.45)
                    tp.paragraph_format.space_after = Pt(6)
                    run = tp.add_run(paragraph.translation)
                    _set_run_font(run, size=9.5, italic=True, color=GRAY)

        _add_quality(document, context)
        _add_footer(document, context)
        document.core_properties.title = context.video.title
        document.core_properties.comments = "Generated by dsh-bilibili-subtitles"
        document.core_properties.author = "dsh-bilibili-subtitles"
        document.save(str(out))

        # Re-open to make sure the package is valid and content exists.
        check = Document(str(out))
        if not check.paragraphs:
            raise RenderError("DOCX 写入后没有段落。")
    except RenderError:
        raise
    except Exception as exc:  # noqa: BLE001
        raise RenderError(f"生成 DOCX 失败: {exc}") from exc
    return out
